"""Recompute/fits for the SYNTHETIC KPI case; elapsed times are NOT service times.

Read-only workbook access. Run from any directory with Python 3.12 and requirements.
Model selection: smallest last-20%-chronological CDF RMSE, including empirical.
This split is used for selection, so it is not independent final validation.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import scipy
from openpyxl import load_workbook
from scipy import stats

ROOT = Path(__file__).resolve().parents[1]
LEGACY_SOURCE = Path("C:/Users/USER/Documents/job-search-ops/applications/2026-10-21-"
                      "career-packs/evidence/inventory-leadtime-kpi-case/"
                      "Inventory_Health_and_LeadTime_KPI_Case.xlsx")
PACKAGED_SOURCE = ROOT / 'analysis/calibration/source-synthetic.xlsx'
DEFAULT_SOURCE = PACKAGED_SOURCE if PACKAGED_SOURCE.exists() else LEGACY_SOURCE
STEPS = ["Kitting", "Assembly", "Test", "Quality release", "Handover"]
KEYS = ["kitting", "assembly", "test", "quality_release", "handover"]
TIMES = ["release", *[f"{k}_end" for k in KEYS]]
HOURS = [f"{k}_elapsed_h" for k in KEYS]
FAMILIES = {"gamma": stats.gamma, "lognormal": stats.lognorm,
            "weibull": stats.weibull_min}
LABEL = "SYNTHETIC - learning data; elapsed includes waiting; not company data"
SEED = 20261023


def read_source(source: Path):
    """Match Excel M6:M205: missing first, otherwise any negative, whole-row exclusion."""
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = wb["LeadTime_Analysis"]
        headers = [ws.cell(5, c).value for c in range(1, 20)]
        rows = list(ws.iter_rows(min_row=6, max_row=205, max_col=13, values_only=True))
        raw = pd.DataFrame([r[:7] for r in rows], columns=["order_id", *TIMES])
        raw.insert(1, "source_row", np.arange(6, 206))
        for col in TIMES:
            raw[col] = pd.to_datetime(raw[col])
        for k, col in enumerate(HOURS):
            raw[col] = (raw[TIMES[k + 1]] - raw[TIMES[k]]).dt.total_seconds() / 3600
        missing = raw[TIMES].isna().any(axis=1)
        negative = (~missing) & raw[HOURS].lt(0).any(axis=1)
        raw["quality_flag"] = np.where(missing, "Missing timestamp",
                                        np.where(negative, "Negative duration", "OK"))
        cached = [r[12] for r in rows]
        if any(flag != actual for flag, actual in zip(raw.quality_flag, cached)):
            raise ValueError("Recomputed quality flags disagree with cached workbook flags")
        expected = {**{f"{key}_mean_h": ws.cell(6 + i, 22).value
                      for i, key in enumerate(KEYS)},
                    **{f"{key}_p90_h": ws.cell(6 + i, 23).value
                       for i, key in enumerate(KEYS)},
                    "end_to_end_mean_h": ws["V11"].value,
                    "end_to_end_p90_h": ws["W11"].value}
        limitation = str(wb["README"]["B13"].value)
    finally:
        wb.close()
    # Keep original Excel row provenance while making all downstream order-based
    # diagnostics chronological, including the IDs assigned to arrival gaps.
    raw = raw.sort_values(["release", "source_row"]).reset_index(drop=True)
    raw["data_label"] = LABEL
    clean = raw[raw.quality_flag == "OK"].copy()
    clean["end_to_end_elapsed_h"] = clean[HOURS].sum(axis=1)
    # Compact numerical aliases for the standalone AnyLogic historical replay.
    clean["release_h"] = (clean.release - raw.release.min()).dt.total_seconds()/3600
    for alias, source_col in zip(["kitting_h", "assembly_h", "test_h", "quality_h", "handover_h"], HOURS):
        clean[alias] = clean[source_col]
    if not np.allclose(clean.end_to_end_elapsed_h,
                       (clean.handover_end-clean.release).dt.total_seconds()/3600):
        raise AssertionError("Telescoping timestamps failed")
    if not (clean[HOURS] > 0).all().all():
        raise ValueError("Positive-support fits require strictly positive durations")
    return raw, clean, headers, expected, limitation


def fit_family(x, family):
    return np.asarray(FAMILIES[family].fit(x, floc=0), dtype=float)


def cdf_at(x, values, family, params):
    if family == "empirical":
        return np.searchsorted(np.sort(x), values, side="right") / len(x)
    return FAMILIES[family].cdf(values, *params)


def sample_model(x, family, params, n, rng):
    if family == "empirical":
        return rng.choice(x, size=n, replace=True)
    return FAMILIES[family].rvs(*params, size=n, random_state=rng)


def bootstrap_summary(matrix, reps=4000):
    """Resample whole orders jointly: retain dependence between elapsed steps."""
    rng = np.random.default_rng(SEED)
    indices = rng.integers(0, len(matrix), size=(reps, len(matrix)))
    draws = matrix[indices]
    means = draws.mean(axis=1)
    p90 = np.quantile(draws, .9, axis=1)
    return np.quantile(means, [.025, .975], axis=0), np.quantile(p90, [.025, .975], axis=0)


def arrivals(raw):
    release = raw.release.sort_values()
    gaps = release.diff().dt.total_seconds().dropna().to_numpy()/3600
    span = (release.iloc[-1] - release.iloc[0]).total_seconds()/3600
    counts = release.dt.floor("D").value_counts().sort_index()
    days = pd.date_range(release.min().floor("D"), release.max().floor("D"))
    daily = counts.reindex(days, fill_value=0)
    ks = stats.kstest(gaps, "expon", args=(0, float(gaps.mean())))
    return {"n_all_releases": len(release), "n_gaps": len(gaps),
            "first_release": release.iloc[0].isoformat(),
            "last_release": release.iloc[-1].isoformat(), "observed_span_h": span,
            "mean_gap_h": float(gaps.mean()), "sd_gap_h": float(gaps.std(ddof=1)),
            "cv_gap": float(gaps.std(ddof=1)/gaps.mean()),
            "p90_gap_h": float(np.quantile(gaps, .9)), "tied_gaps": int((gaps == 0).sum()),
            "exponential_ks_D_descriptive": float(ks.statistic),
            "rate_inverse_mean_gap_per_h": float(1/gaps.mean()),
            "generator_window_h": 42*24, "generator_window_rate_per_h": 200/(42*24),
            "generator_mechanism": "200 fixed-count independent uniform releases over 42 days, rounded to minutes; not a Poisson process",
            "caution": "Gap exponential compatibility is descriptive; fixed count and finite window induce dependence. Rate from clean195 releases would incorrectly treat quality exclusion as lost demand.",
            "service_proxy_offered_loads": dict(zip(KEYS, (raw.loc[raw.quality_flag == 'OK', HOURS].mean()/gaps.mean()).tolist()))}, gaps, daily


def run(source=DEFAULT_SOURCE, bootstrap_reps=4000):
    out = ROOT / "analysis/calibration"
    results = ROOT / "results/calibration"
    out.mkdir(parents=True, exist_ok=True)
    results.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(source.read_bytes()).hexdigest()
    raw, clean, headers, expected, limitation = read_source(source)
    for name, frame in (("clean_orders", clean), ("rejected_records", raw[raw.quality_flag != "OK"])):
        frame.to_csv(out/f"{name}.csv", index=False)
    matrix = np.column_stack([clean[HOURS].to_numpy(), clean.end_to_end_elapsed_h])
    mean_ci, p90_ci = bootstrap_summary(matrix, bootstrap_reps)
    summary_rows = []
    for i, key in enumerate([*KEYS, "end_to_end"]):
        x = matrix[:, i]
        mean, p90 = float(x.mean()), float(np.quantile(x, .9))
        if not np.isclose(mean, expected[f"{key}_mean_h"], atol=1e-7):
            raise AssertionError(f"Workbook mean mismatch {key}")
        if not np.isclose(p90, expected[f"{key}_p90_h"], atol=1e-7):
            raise AssertionError(f"Workbook P90 mismatch {key}")
        summary_rows.append(dict(data_label=LABEL, step=key, n=len(x), mean_h=mean,
                                 mean_ci_low_h=mean_ci[0, i], mean_ci_high_h=mean_ci[1, i],
                                 p90_h=p90, p90_ci_low_h=p90_ci[0, i],
                                 p90_ci_high_h=p90_ci[1, i]))
    pd.DataFrame(summary_rows).to_csv(results/"observed_summary.csv", index=False)
    records, selected, models = [], [], []
    fig, axes = plt.subplots(2, 3, figsize=(14, 8))
    train_n = int(.8*len(clean))
    for i, key in enumerate(KEYS):
        x = clean[HOURS[i]].to_numpy()
        train, holdout = x[:train_n], x[train_n:]
        values = np.sort(holdout)
        target = (np.arange(len(values))+.5)/len(values)
        station_fits = []
        for family in [*FAMILIES, "empirical"]:
            params = [] if family == "empirical" else fit_family(x, family).tolist()
            train_params = [] if family == "empirical" else fit_family(train, family)
            prediction = cdf_at(train, values, family, train_params)
            rmse = float(np.sqrt(np.mean((prediction-target)**2)))
            if family == "empirical":
                aicc, ks = None, None
                fitted_mean, fitted_p90 = x.mean(), np.quantile(x, .9)
            else:
                dist = FAMILIES[family]
                ll = float(dist.logpdf(x, *params).sum())
                # Location is fixed at zero; only shape and scale are estimated.
                aicc = 4 - 2*ll + 12/(len(x)-3)
                ks = float(stats.kstest(x, dist.cdf, args=tuple(params)).statistic)
                fitted_mean, fitted_p90 = dist.mean(*params), dist.ppf(.9, *params)
            row = dict(data_label=LABEL, step=key, family=family, n=len(x),
                       params_scipy_json=json.dumps(params), aicc=aicc,
                       ks_D_descriptive=ks, chronological_selection_cdf_rmse=rmse,
                       selection_train_n=len(train), selection_test_n=len(holdout),
                       fitted_mean_h=float(fitted_mean), fitted_p90_h=float(fitted_p90),
                       source_columns=f"LeadTime_Analysis!{'BCDEF'[i]}6:{'BCDEF'[i]}205 -> {'CDEFG'[i]}6:{'CDEFG'[i]}205",
                       interpretation="elapsed time including waiting; service proxy only by explicit assumption")
            station_fits.append(row)
            records.append(row)
        winner = min(station_fits, key=lambda d: d["chronological_selection_cdf_rmse"])
        selected.append(winner.copy())
        models.append((winner["family"], json.loads(winner["params_scipy_json"])))
        grid = np.linspace(0, x.max()*1.05, 300)
        ax = axes.flat[i]
        ax.step(np.sort(x), np.arange(1, len(x)+1)/len(x), where="post", label="Observed ECDF", color="black")
        for row in station_fits:
            if row["family"] != "empirical":
                ax.plot(grid, cdf_at(x, grid, row["family"], json.loads(row["params_scipy_json"])),
                        label=row["family"], alpha=.8)
        ax.set(title=STEPS[i], xlabel="Elapsed hours", ylabel="CDF")
        ax.legend(fontsize=7)
    axes.flat[-1].axis("off")
    axes.flat[-1].text(0, .8, "SYNTHETIC\nElapsed includes waiting.\nFitted marginals do not identify\nservice times or capacities.\n\nKS D descriptive only;\nno invalid fitted-parameter p-values.", va="top")
    fig.suptitle("SYNTHETIC elapsed-time fits: empirical, gamma, lognormal, Weibull")
    fig.tight_layout()
    fig.savefig(results/"fit_cdfs.png", dpi=160)
    plt.close(fig)
    pd.DataFrame(records).to_csv(out/"fit_diagnostics.csv", index=False)
    pd.DataFrame(selected).to_csv(out/"calibration_table.csv", index=False)
    corr = clean[HOURS].corr(method="spearman")
    corr.index = KEYS
    corr.columns = KEYS
    corr.to_csv(results/"step_spearman_correlations.csv", index_label="SYNTHETIC_step")
    ar, gaps, daily = arrivals(raw)
    pd.DataFrame({"data_label": LABEL, "day": daily.index, "releases": daily.to_numpy()}).to_csv(results/"daily_releases.csv", index=False)
    gap_frame = pd.DataFrame({"data_label": LABEL, "preceding_order_id": raw.order_id.iloc[:-1].to_numpy(),
                              "order_id": raw.order_id.iloc[1:].to_numpy(), "interarrival_h": gaps})
    gap_frame.to_csv(out/"arrival_gaps.csv", index=False)
    # This is an input-model reproduction diagnostic, not a capacity simulation.
    rng = np.random.default_rng(SEED+1)
    independent = np.column_stack([sample_model(clean[HOURS[i]].to_numpy(), *models[i], 100000, rng) for i in range(5)])
    joint = matrix[rng.integers(0, len(matrix), size=100000), :5]
    compare = []
    for name, arr in (("selected_independent_marginals", independent), ("joint_empirical_orders", joint)):
        full = np.column_stack([arr, arr.sum(axis=1)])
        for i, key in enumerate([*KEYS, "end_to_end"]):
            compare.append(dict(data_label=LABEL, input_model=name, step=key, n=len(arr),
                                reproduced_mean_h=float(full[:, i].mean()),
                                reproduced_p90_h=float(np.quantile(full[:, i], .9)),
                                observed_mean_h=summary_rows[i]["mean_h"],
                                observed_p90_h=summary_rows[i]["p90_h"],
                                status="input-model reproduction only; no AnyLogic run or capacity validation"))
    pd.DataFrame(compare).to_csv(results/"input_model_comparison.csv", index=False)
    fig, ax = plt.subplots(1, 2, figsize=(11, 4))
    ax[0].hist(matrix[:, -1], bins=18, density=True, alpha=.6, label="Observed195")
    ax[0].hist(independent.sum(axis=1), bins=45, density=True, histtype="step", label="Independent fitted inputs")
    ax[0].set(xlabel="End-to-end elapsed hours", ylabel="Density")
    ax[0].legend(fontsize=8)
    ax[1].hist(gaps, bins=20, density=True, alpha=.6, label="Observed199 gaps")
    grid = np.linspace(0, gaps.max(), 300)
    ax[1].plot(grid, stats.expon.pdf(grid, scale=gaps.mean()), label="Exponential reference only")
    ax[1].set(xlabel="Release gap hours", ylabel="Density")
    ax[1].legend(fontsize=8)
    fig.suptitle("SYNTHETIC elapsed reconstruction and arrivals; not capacity validation")
    fig.tight_layout()
    fig.savefig(results/"elapsed_and_arrivals.png", dpi=160)
    plt.close(fig)
    payload = dict(data_label=LABEL, source=str(source.resolve()), source_sha256=digest,
                   source_sheet="LeadTime_Analysis", source_rows="6:205", source_headers=headers,
                   counts=raw.quality_flag.value_counts().to_dict(), raw_orders=len(raw), clean_orders=len(clean),
                   excluded_ids=raw.loc[raw.quality_flag != "OK", ["order_id", "source_row", "quality_flag"]].to_dict("records"),
                   observed=summary_rows, arrivals=ar,
                   selection="Minimum CDF RMSE on chronological final39 versus first156. Selection uses this holdout; not independent final validation.",
                   bootstrap=dict(method="whole-order nonparametric percentile bootstrap; 95% intervals for descriptive mean/P90", reps=bootstrap_reps, seed=SEED,
                                  limitations="assumes orders iid; does not include structural/service decomposition uncertainty"),
                   source_limitation=limitation, time_unit="hour", service_times_identifiable=False,
                   capacities_identifiable=False, rework_identifiable=False,
                   dependency=dict(spearman_max_abs_off_diagonal=float(np.abs(corr.to_numpy()-np.eye(5)).max()),
                                   rule="Joint row sampling preserves observed cross-step dependence; independent fitted marginals assume it away. Small empirical correlations do not establish independence."),
                   caveats=["Only completion timestamps observed, so elapsed = waiting + processing; resource capacities, utilization, queues and starvation cannot be recovered.",
                            "Generator uses independent positive lognormal durations, but source documentation labels them elapsed including waiting; do not reinterpret as observed service.",
                            "Workbook data excludes incomplete/negative whole orders; retain all200 releases for demand. Observed timing precision one minute.",
                            "Kitting component quantities, processing starts, station resources, shifts, test failures/rework and costs are absent.",
                            "Adding a capacity station can only be hypothetical under labelled assumptions; original data cannot validate the causal benefit."],
                   versions=dict(numpy=np.__version__, pandas=pd.__version__, scipy=scipy.__version__),
                   source_unchanged=hashlib.sha256(source.read_bytes()).hexdigest() == digest)
    if not payload["source_unchanged"]:
        raise AssertionError("Source hash changed")
    (out/"summary.json").write_text(json.dumps(payload, indent=2, allow_nan=False), encoding="utf-8")
    print(json.dumps({"data_label": LABEL, "raw": len(raw), "clean": len(clean),
                      "mean_e2e_h": matrix[:, -1].mean(), "p90_e2e_h": np.quantile(matrix[:, -1], .9),
                      "source_unchanged": True, "selected": {r['step']: r['family'] for r in selected}}, indent=2))
    return payload


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--bootstrap-reps", type=int, default=4000)
    args = parser.parse_args()
    if args.bootstrap_reps < 100:
        parser.error("--bootstrap-reps must be at least100")
    run(args.source, args.bootstrap_reps)
